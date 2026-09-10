from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.http import Http404, HttpResponse
from django.core.paginator import Paginator
from django.shortcuts import redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .forms import (
    ClienteForm,
    ClienteQiBuscaForm,
    ComercialAnalysisUploadForm,
    ContatoForm,
    FundoBuscaForm,
    FundoCatalogoBuscaForm,
    FundoReceitaForm,
    RegraReceitaForm,
    UploadCarteiraForm,
)
from .services.cvm_sync import (
    discover_qi_client_candidates,
    get_sync_metadata,
    import_qi_client_candidate,
    list_catalog_funds,
    search_funds,
    sync_cvm_data,
)
from .services.commercial_analysis import (
    analysis_files_available,
    build_and_cache_analysis_snapshot,
    get_analysis_files,
    load_cached_analysis_snapshot,
    save_uploaded_analysis_files,
)
from .services.data_store import (
    add_client,
    add_contact,
    add_fund,
    build_dashboard_metrics,
    build_user_dashboard,
    delete_client,
    delete_rule,
    ensure_storage,
    format_currency,
    get_client,
    get_client_by_name,
    group_clients_by_category,
    list_category_choices,
    list_clients_overview,
    list_known_product_types,
    list_rules,
    parse_decimal,
    refresh_client_revenue,
    set_client_user_status,
    set_fund_user_status,
    update_client_category,
    update_client_company_identity,
    upsert_company_profile,
    upsert_rule,
    update_fund_fee_rate,
)
from .services.importers import import_clients_from_file
from .services.public_company import enrich_company_profiles_for_client


def _matches_text_filter(value: str, operator: str, filter_value: str) -> bool:
    normalized_value = " ".join(str(value or "").strip().lower().split())
    normalized_filter = " ".join(str(filter_value or "").strip().lower().split())
    if not normalized_filter:
        return True
    if operator == "eq":
        return normalized_value == normalized_filter
    if operator == "neq":
        return normalized_value != normalized_filter
    return normalized_filter in normalized_value


def dashboard(request):
    ensure_storage()
    context = {
        "metrics": build_dashboard_metrics(),
        "groups": group_clients_by_category(),
        "category_choices": list_category_choices(),
        "cvm_metadata": get_sync_metadata(),
    }
    return render(request, "carteira/dashboard.html", context)


def dashboard_usuario(request):
    ensure_storage()
    filters = {
        "fund_type": request.GET.get("fund_type", ""),
        "manager_name": request.GET.get("manager_name", ""),
        "administrator_name": request.GET.get("administrator_name", ""),
        "category": request.GET.get("category", ""),
    }
    context = build_user_dashboard(filters)
    return render(request, "carteira/dashboard_usuario.html", context)


def analise_comercial(request):
    ensure_storage()
    summary = None
    error_message = None

    if request.method == "POST":
        form = ComercialAnalysisUploadForm(request.POST, request.FILES)
        if form.is_valid():
            files = get_analysis_files()
            scorecard_file = form.cleaned_data["scorecard"]
            base_file = form.cleaned_data["base_fundos"]
            if (scorecard_file and not base_file and not files.base_path.exists()) or (
                base_file and not scorecard_file and not files.scorecard_path.exists()
            ):
                error_message = "Na primeira carga, envie as duas planilhas para montar a analise completa."
            else:
                save_uploaded_analysis_files(scorecard_file, base_file)
                try:
                    summary = build_and_cache_analysis_snapshot()
                except Exception as exc:
                    error_message = f"Nao foi possivel processar as planilhas: {exc}"
        else:
            error_message = "Revise os arquivos enviados e tente novamente."
    else:
        form = ComercialAnalysisUploadForm()

    if summary is None and not error_message and analysis_files_available():
        try:
            summary = load_cached_analysis_snapshot() or build_and_cache_analysis_snapshot()
        except Exception as exc:
            error_message = f"Nao foi possivel carregar a analise salva: {exc}"

    return render(
        request,
        "carteira/analise_comercial.html",
        {
            "form": form,
            "summary": summary,
            "error_message": error_message,
            "has_saved_files": analysis_files_available(),
        },
    )


def lista_clientes(request):
    ensure_storage()
    search_term = (request.GET.get("consulta") or "").strip()
    context = {
        "clients": list_clients_overview(search_term),
        "search_term": search_term,
    }
    return render(request, "carteira/clientes.html", context)


def alternar_cliente_usuario(request, client_id: str):
    ensure_storage()
    client = get_client(client_id)
    if not client:
        raise Http404("Cliente nao encontrado.")
    next_url = request.POST.get("next") or request.GET.get("next")
    if request.method == "POST":
        set_client_user_status(client_id, request.POST.get("is_user_client", "0") == "1")
    if next_url:
        return redirect(next_url)
    return redirect("lista_clientes")


def atualizar_categoria_cliente(request, client_id: str):
    ensure_storage()
    client = get_client(client_id)
    if not client:
        raise Http404("Cliente nao encontrado.")
    if request.method == "POST":
        update_client_category(client_id, request.POST.get("category", ""))
    return redirect("dashboard")


def importar_carteira(request):
    ensure_storage()
    summary = None

    if request.method == "POST":
        form = UploadCarteiraForm(request.POST, request.FILES)
        if form.is_valid():
            upload = form.cleaned_data["arquivo"]
            destination = Path(settings.UPLOAD_DIR) / upload.name
            with destination.open("wb+") as file_handle:
                for chunk in upload.chunks():
                    file_handle.write(chunk)
            summary = import_clients_from_file(destination)
    else:
        form = UploadCarteiraForm()

    return render(
        request,
        "carteira/importar.html",
        {
            "form": form,
            "summary": summary,
        },
    )


def novo_cliente(request):
    ensure_storage()
    if request.method == "POST":
        form = ClienteForm(request.POST)
        if form.is_valid():
            payload = {
                "name": form.cleaned_data["nome"],
                "category": form.cleaned_data["categoria"],
                "monthly_revenue": str(form.cleaned_data["receita_mensal"] or ""),
                "monthly_revenue_source": form.cleaned_data["fonte_receita"],
                "notes": form.cleaned_data["observacoes"],
            }
            client = add_client(payload)
            return redirect("detalhe_cliente", client_id=client["id"])
    else:
        form = ClienteForm()

    return render(request, "carteira/novo_cliente.html", {"form": form})


def detalhe_cliente(request, client_id: str):
    ensure_storage()
    client = get_client(client_id)
    if not client:
        raise Http404("Cliente nao encontrado.")
    return render(
        request,
        "carteira/detalhe_cliente.html",
        {
            "client": client,
            "cvm_metadata": get_sync_metadata(),
            "contact_form": ContatoForm(),
            "fund_type_options": list_known_product_types(),
        },
    )


def sincronizar_cvm(request):
    ensure_storage()
    sync_result = None
    error_message = None
    if request.method == "POST":
        try:
            sync_result = sync_cvm_data()
        except Exception as exc:
            error_message = str(exc)

    return render(
        request,
        "carteira/cvm_sync.html",
        {
            "sync_result": sync_result,
            "error_message": error_message,
            "metadata": get_sync_metadata(),
        },
    )


def buscar_fundos_cliente(request, client_id: str):
    ensure_storage()
    client = get_client(client_id)
    if not client:
        raise Http404("Cliente nao encontrado.")

    metadata = get_sync_metadata()
    results: list[dict[str, str]] = []
    if request.GET:
        form = FundoBuscaForm(request.GET)
        if form.is_valid() and metadata:
            results = search_funds(
                search_term=form.cleaned_data["consulta"] or client["name"],
                qitech_only=form.cleaned_data["somente_qi"],
            )
    else:
        form = FundoBuscaForm(initial={"consulta": client["name"], "somente_qi": True})
        if metadata:
            results = search_funds(search_term=client["name"], qitech_only=True)

    return render(
        request,
        "carteira/buscar_fundos.html",
        {
            "client": client,
            "form": form,
            "results": results,
            "metadata": metadata,
        },
    )


def vincular_fundo_cliente(request, client_id: str):
    ensure_storage()
    client = get_client(client_id)
    if not client:
        raise Http404("Cliente nao encontrado.")
    if request.method != "POST":
        return redirect("buscar_fundos_cliente", client_id=client_id)

    payload = {
        "client_id": client_id,
        "cvm_fund_id": request.POST.get("cvm_fund_id", ""),
        "cvm_class_id": request.POST.get("cvm_class_id", ""),
        "fund_name": request.POST.get("fund_name", ""),
        "cnpj": request.POST.get("cnpj", ""),
        "pl": request.POST.get("pl", ""),
        "monthly_revenue": request.POST.get("monthly_revenue", ""),
        "product_type": request.POST.get("product_type", ""),
        "manager_name": request.POST.get("manager_name", ""),
        "administrator_name": request.POST.get("administrator_name", ""),
        "status": request.POST.get("status", ""),
        "regulation_url": request.POST.get("regulation_url", ""),
        "revenue_source": "CVM Dados Abertos",
    }
    if payload["fund_name"] and payload["cnpj"]:
        add_fund(payload)
    return redirect("detalhe_cliente", client_id=client_id)


def regras_receita(request):
    ensure_storage()
    if request.method == "POST":
        form = RegraReceitaForm(request.POST)
        if form.is_valid():
            upsert_rule(
                product_type=form.cleaned_data["product_type"],
                annual_fee_rate=str(form.cleaned_data["annual_fee_rate"]),
                notes=form.cleaned_data["notes"],
            )
            return redirect("regras_receita")
    else:
        form = RegraReceitaForm()

    return render(
        request,
        "carteira/regras_receita.html",
        {
            "form": form,
            "rules": list_rules(),
            "known_product_types": list_known_product_types(),
        },
    )


def remover_regra_receita(request, product_type: str):
    ensure_storage()
    if request.method == "POST":
        delete_rule(product_type)
    return redirect("regras_receita")


def atualizar_receita_fundo(request, client_id: str):
    ensure_storage()
    client = get_client(client_id)
    if not client:
        raise Http404("Cliente nao encontrado.")
    if request.method != "POST":
        return redirect("detalhe_cliente", client_id=client_id)

    form = FundoReceitaForm(request.POST)
    if form.is_valid():
        update_fund_fee_rate(
            fund_id=form.cleaned_data["fund_id"],
            annual_fee_rate=str(form.cleaned_data["annual_fee_rate"]),
        )
        refresh_client_revenue(client_id)
    return redirect("detalhe_cliente", client_id=client_id)


def adicionar_contato_cliente(request, client_id: str):
    ensure_storage()
    client = get_client(client_id)
    if not client:
        raise Http404("Cliente nao encontrado.")
    if request.method != "POST":
        return redirect("detalhe_cliente", client_id=client_id)

    form = ContatoForm(request.POST)
    if form.is_valid():
        add_contact(
            {
                "client_id": client_id,
                "name": form.cleaned_data["nome"],
                "role": form.cleaned_data["cargo"],
                "area": form.cleaned_data["area"],
                "email": form.cleaned_data["email"],
                "phone": form.cleaned_data["telefone"],
                "linkedin": form.cleaned_data["linkedin"],
                "notes": form.cleaned_data["observacoes"],
            }
        )
    return redirect("detalhe_cliente", client_id=client_id)


def enriquecer_cliente(request, client_id: str):
    ensure_storage()
    client = get_client(client_id)
    if not client:
        raise Http404("Cliente nao encontrado.")
    if request.method != "POST":
        return redirect("detalhe_cliente", client_id=client_id)

    if not str(client.get("company_cnpj", "")).strip():
        candidate = next(
            (
                item
                for item in discover_qi_client_candidates(limit=1000)
                if str(item.get("name", "")).strip().lower() == str(client.get("name", "")).strip().lower()
            ),
            None,
        )
        if candidate and candidate.get("company_cnpj"):
            update_client_company_identity(
                client_id,
                str(candidate.get("company_cnpj", "")),
                str(candidate.get("company_cnpj_source", "")),
            )
            client = get_client(client_id)

    profiles = enrich_company_profiles_for_client(client)
    for profile in profiles:
        upsert_company_profile(profile)
    return redirect("detalhe_cliente", client_id=client_id)


def alternar_fundo_usuario(request, client_id: str, fund_id: str):
    ensure_storage()
    client = get_client(client_id)
    if not client:
        raise Http404("Cliente nao encontrado.")
    next_url = request.POST.get("next") or request.GET.get("next")
    if request.method == "POST":
        set_fund_user_status(fund_id, request.POST.get("is_user_fund", "0") == "1")
    if next_url:
        return redirect(next_url)
    return redirect("detalhe_cliente", client_id=client_id)


def remover_cliente(request, client_id: str):
    ensure_storage()
    client = get_client(client_id)
    if not client:
        raise Http404("Cliente nao encontrado.")
    if request.method == "POST":
        delete_client(client_id)
        return redirect("dashboard")
    return redirect("detalhe_cliente", client_id=client_id)


def clientes_qi_descobertos(request):
    ensure_storage()
    metadata = get_sync_metadata()
    import_summary = None

    if request.method == "POST" and metadata:
        action = request.POST.get("action", "")
        if action == "import_one":
            client_name = request.POST.get("client_name", "")
            exists_before = get_client_by_name(client_name) is not None
            result = import_qi_client_candidate(client_name)
            if result:
                import_summary = {
                    "imported_clients": 0 if exists_before else 1,
                    "linked_funds": int(result["fund_count"]),
                    "single_client": client_name,
                }
        elif action == "import_selected":
            selected_client_names = []
            seen_names: set[str] = set()
            for client_name in request.POST.getlist("client_names"):
                normalized_name = client_name.strip()
                if not normalized_name or normalized_name in seen_names:
                    continue
                seen_names.add(normalized_name)
                selected_client_names.append(normalized_name)

            imported_clients = 0
            linked_funds = 0
            processed_clients = 0
            for client_name in selected_client_names:
                exists_before = get_client_by_name(client_name) is not None
                result = import_qi_client_candidate(client_name)
                if not result:
                    continue
                processed_clients += 1
                linked_funds += int(result["fund_count"])
                if not exists_before:
                    imported_clients += 1

            if processed_clients:
                import_summary = {
                    "imported_clients": imported_clients,
                    "linked_funds": linked_funds,
                    "selected_count": processed_clients,
                }

    search_form = ClienteQiBuscaForm(request.GET or None)
    candidates = discover_qi_client_candidates() if metadata else []
    if metadata and search_form.is_valid():
        query = (search_form.cleaned_data.get("consulta") or "").strip().lower()
        if query:
            candidates = [
                candidate
                for candidate in candidates
                if query in str(candidate.get("name", "")).lower()
                or query in str(candidate.get("company_cnpj", "")).lower()
                or query in str(candidate.get("category", "")).lower()
                or any(query in str(fund.get("fund_name", "")).lower() for fund in candidate.get("funds", []))
            ]

    return render(
        request,
        "carteira/clientes_qi_descobertos.html",
        {
            "metadata": metadata,
            "candidates": candidates,
            "import_summary": import_summary,
            "search_form": search_form,
        },
    )


def catalogo_fundos(request):
    ensure_storage()
    search_form = FundoCatalogoBuscaForm(request.GET or None)
    funds, catalog_filters, fund_type_options = _get_filtered_catalog_funds(request, search_form)
    total_pl = sum((parse_decimal(fund.get("pl", "0")) for fund in funds), start=parse_decimal("0"))
    per_page = request.GET.get("per_page") or "20"

    per_page_value = int(per_page) if per_page in {"20", "50", "100"} else 20
    paginator = Paginator(funds, per_page_value)
    page_obj = paginator.get_page(request.GET.get("page"))
    pagination_params = request.GET.copy()
    pagination_params.pop("page", None)
    sync_metadata = get_sync_metadata() or {}
    reference_dates = [str(fund.get("pl_date", "")) for fund in funds if fund.get("pl_date")]
    context = {
        "search_form": search_form,
        "page_obj": page_obj,
        "pagination_range": paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1),
        "pagination_ellipsis": paginator.ELLIPSIS,
        "pagination_query": pagination_params.urlencode(),
        "funds": page_obj.object_list,
        "fund_count": len(funds),
        "total_pl": format_currency(total_pl),
        "last_cvm_sync": _format_datetime(sync_metadata.get("imported_at", "")),
        "latest_pl_reference": _format_date(max(reference_dates, default="")),
        "active_filters": _build_active_filter_chips(request),
        "fund_type_options": fund_type_options,
        **catalog_filters,
        "per_page": str(per_page_value),
    }
    return render(request, "carteira/catalogo_fundos.html", context)


def _get_filtered_catalog_funds(request, search_form):
    search_term = ""
    if search_form.is_valid():
        search_term = search_form.cleaned_data.get("consulta") or ""
    funds = list_catalog_funds(search_term=search_term, qitech_only=False)
    fund_type_filter = (request.GET.get("fund_type") or "").strip()
    manager_operator = (request.GET.get("manager_operator") or "contains").strip()
    manager_filter = (request.GET.get("manager_filter") or "").strip().lower()
    administrator_operator = (request.GET.get("administrator_operator") or "contains").strip()
    administrator_filter = (request.GET.get("administrator_filter") or "").strip().lower()
    sort = (request.GET.get("sort") or "").strip()
    fund_type_options = sorted({fund.get("product_type", "") for fund in funds if fund.get("product_type", "")})

    if fund_type_filter:
        funds = [fund for fund in funds if fund.get("product_type", "") == fund_type_filter]
    if manager_filter:
        funds = [
            fund
            for fund in funds
            if _matches_text_filter(str(fund.get("manager_name", "")), manager_operator, manager_filter)
        ]
    if administrator_filter:
        funds = [
            fund
            for fund in funds
            if _matches_text_filter(
                str(fund.get("administrator_name", "")),
                administrator_operator,
                administrator_filter,
            )
        ]

    if sort == "pl_asc":
        funds = sorted(funds, key=lambda item: parse_decimal(item.get("pl", "0")))
    elif sort == "pl_desc":
        funds = sorted(funds, key=lambda item: parse_decimal(item.get("pl", "0")), reverse=True)
    elif sort == "revenue_asc":
        funds = sorted(funds, key=lambda item: parse_decimal(item.get("monthly_revenue", "0")))
    elif sort == "revenue_desc":
        funds = sorted(funds, key=lambda item: parse_decimal(item.get("monthly_revenue", "0")), reverse=True)
    elif sort == "name":
        funds = sorted(funds, key=lambda item: str(item.get("fund_name", "")).casefold())

    filters = {
        "fund_type_filter": fund_type_filter,
        "manager_operator": manager_operator,
        "manager_filter": manager_filter,
        "administrator_operator": administrator_operator,
        "administrator_filter": administrator_filter,
        "sort": sort,
    }
    return funds, filters, fund_type_options


def exportar_fundos_excel(request):
    ensure_storage()
    search_form = FundoCatalogoBuscaForm(request.GET or None)
    funds, _, _ = _get_filtered_catalog_funds(request, search_form)

    column_definitions = {
        "fund_name": ("Fundo", "fund_name", 52),
        "cnpj": ("CNPJ", "cnpj", 20),
        "pl": ("PL", "pl", 20),
        "pl_date": ("Data de referencia do PL", "pl_date", 25),
        "product_type": ("Tipo de fundo", "product_type", 18),
        "manager_name": ("Gestora", "manager_name", 38),
        "administrator_name": ("Administrador", "administrator_name", 38),
        "status": ("Status", "status", 28),
    }
    selected_columns = [
        column for column in request.GET.getlist("columns") if column in column_definitions
    ] or list(column_definitions)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Fundos"
    worksheet.append([column_definitions[column][0] for column in selected_columns])

    header_fill = PatternFill("solid", fgColor="0F1E3E")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")

    for fund in funds:
        row = []
        for column in selected_columns:
            value = fund.get(column_definitions[column][1], "")
            if column == "pl":
                value = float(parse_decimal(value))
            elif column == "pl_date":
                value = parse_date(str(value)) if value else None
            row.append(value)
        worksheet.append(row)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, column in enumerate(selected_columns, start=1):
        worksheet.column_dimensions[worksheet.cell(row=1, column=index).column_letter].width = column_definitions[column][2]
        if column == "pl":
            for cell in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for item in cell:
                    item.number_format = 'R$ #,##0.00'
        elif column == "pl_date":
            for cell in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for item in cell:
                    item.number_format = "dd/mm/yyyy"

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="catalogo-de-fundos.xlsx"'
    return response


def _format_datetime(value: object) -> str:
    parsed = parse_datetime(str(value or ""))
    if not parsed:
        return str(value or "")
    if timezone.is_aware(parsed):
        parsed = timezone.localtime(parsed)
    return parsed.strftime("%d/%m/%Y às %H:%M")


def _format_date(value: object) -> str:
    parsed = parse_date(str(value or ""))
    return parsed.strftime("%d/%m/%Y") if parsed else str(value or "")


def _query_without(request, *keys: str) -> str:
    query = request.GET.copy()
    query.pop("page", None)
    for key in keys:
        query.pop(key, None)
    encoded = query.urlencode()
    return f"{request.path}?{encoded}" if encoded else request.path


def _build_active_filter_chips(request) -> list[dict[str, str]]:
    chips = []
    search_term = (request.GET.get("consulta") or "").strip()
    if search_term:
        chips.append({"label": f"Busca: {search_term}", "remove_url": _query_without(request, "consulta")})

    fund_type = (request.GET.get("fund_type") or "").strip()
    if fund_type:
        chips.append({"label": f"Tipo: {fund_type}", "remove_url": _query_without(request, "fund_type")})

    operator_labels = {"contains": "contém", "eq": "igual a", "neq": "diferente de"}
    manager = (request.GET.get("manager_filter") or "").strip()
    if manager:
        operator = operator_labels.get(request.GET.get("manager_operator", "contains"), "contém")
        chips.append(
            {
                "label": f"Gestora {operator}: {manager}",
                "remove_url": _query_without(request, "manager_filter", "manager_operator"),
            }
        )

    administrator = (request.GET.get("administrator_filter") or "").strip()
    if administrator:
        operator = operator_labels.get(request.GET.get("administrator_operator", "contains"), "contém")
        chips.append(
            {
                "label": f"Administrador {operator}: {administrator}",
                "remove_url": _query_without(request, "administrator_filter", "administrator_operator"),
            }
        )

    sort_labels = {
        "name": "Nome do fundo",
        "pl_asc": "PL crescente",
        "pl_desc": "PL decrescente",
        "revenue_asc": "Receita crescente",
        "revenue_desc": "Receita decrescente",
    }
    sort = (request.GET.get("sort") or "").strip()
    if sort in sort_labels:
        chips.append({"label": f"Ordem: {sort_labels[sort]}", "remove_url": _query_without(request, "sort")})
    return chips
