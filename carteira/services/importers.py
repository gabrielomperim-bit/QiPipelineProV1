import csv
from pathlib import Path

from openpyxl import load_workbook

from .data_store import ImportSummary, add_client

CATEGORY_ALIASES = {
    "gestora": "gestora",
    "consultoria": "consultoria",
    "banco": "banco",
    "securitizadora": "securitizadora",
    "securitizadora ": "securitizadora",
    "outros": "outros",
}

NAME_HEADERS = {"cliente", "nome", "empresa", "conta"}
CATEGORY_HEADERS = {"categoria", "tipo", "classificacao", "classificaÃ§Ã£o"}
REVENUE_HEADERS = {"receita_mensal", "receita mensal", "receita"}
SOURCE_HEADERS = {"fonte_receita", "fonte", "origem_receita"}
NOTES_HEADERS = {"observacoes", "observaÃ§Ãµes", "notas"}


def import_clients_from_file(path: Path) -> ImportSummary:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _read_csv(path)
    elif suffix == ".xlsx":
        rows = _read_xlsx(path)
    else:
        return ImportSummary(
            imported_count=0,
            skipped_count=1,
            details=["Formato nao suportado. Use .csv ou .xlsx."],
        )

    normalized = _normalize_rows(rows)
    imported_count = 0
    skipped_count = 0
    details: list[str] = []

    for row in normalized:
        if not row.get("name"):
            skipped_count += 1
            details.append("Linha ignorada porque o nome do cliente veio vazio.")
            continue
        add_client(row)
        imported_count += 1

    if imported_count:
        details.append(f"{imported_count} clientes importados com sucesso.")

    return ImportSummary(
        imported_count=imported_count,
        skipped_count=skipped_count,
        details=details,
    )


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file_handle:
        return list(csv.DictReader(file_handle))


def _read_xlsx(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(values_only=True))]
    rows: list[dict[str, str]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = {}
        for header, value in zip(headers, values):
            row[header] = "" if value is None else str(value).strip()
        rows.append(row)
    return rows


def _normalize_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    if not rows:
        return []

    headers = {_normalize_header(key) for key in rows[0].keys()}
    category_matrix_mode = any(header in CATEGORY_ALIASES for header in headers) and not (
        headers & NAME_HEADERS
    )

    if category_matrix_mode:
        return _normalize_matrix_rows(rows)
    return _normalize_standard_rows(rows)


def _normalize_matrix_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    normalized: list[dict[str, str]] = []
    for row in rows:
        for original_key, value in row.items():
            key = _normalize_header(original_key)
            if key not in CATEGORY_ALIASES or not str(value).strip():
                continue
            normalized.append(
                {
                    "name": str(value).strip(),
                    "category": CATEGORY_ALIASES[key],
                    "monthly_revenue": "",
                    "monthly_revenue_source": "",
                    "notes": "",
                }
            )
    return normalized


def _normalize_standard_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    normalized: list[dict[str, str]] = []
    for row in rows:
        mapped = {_normalize_header(key): str(value).strip() for key, value in row.items()}
        category = _find_first_value(mapped, CATEGORY_HEADERS)
        normalized.append(
            {
                "name": _find_first_value(mapped, NAME_HEADERS),
                "category": CATEGORY_ALIASES.get(category.lower(), "outros") if category else "outros",
                "monthly_revenue": _find_first_value(mapped, REVENUE_HEADERS),
                "monthly_revenue_source": _find_first_value(mapped, SOURCE_HEADERS),
                "notes": _find_first_value(mapped, NOTES_HEADERS),
            }
        )
    return normalized


def _find_first_value(mapped: dict[str, str], options: set[str]) -> str:
    for option in options:
        if option in mapped and mapped[option]:
            return mapped[option]
    return ""


def _normalize_header(value: str) -> str:
    return str(value).strip().lower().replace("-", "_")
