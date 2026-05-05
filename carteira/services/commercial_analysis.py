from __future__ import annotations

import json
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Iterable

from django.conf import settings
from openpyxl import load_workbook

from .data_store import format_currency

SCORECARD_UPLOAD_NAME = "analise_scorecard.xlsx"
BASE_UPLOAD_NAME = "analise_base.xlsx"
SNAPSHOT_NAME = "analise_comercial_snapshot.json"
METADATA_NAME = "analise_comercial_metadata.json"


@dataclass
class AnalysisFiles:
    scorecard_path: Path
    base_path: Path
    snapshot_path: Path
    metadata_path: Path


def get_analysis_files() -> AnalysisFiles:
    upload_dir = Path(settings.UPLOAD_DIR)
    return AnalysisFiles(
        scorecard_path=upload_dir / SCORECARD_UPLOAD_NAME,
        base_path=upload_dir / BASE_UPLOAD_NAME,
        snapshot_path=upload_dir / SNAPSHOT_NAME,
        metadata_path=upload_dir / METADATA_NAME,
    )


def save_uploaded_analysis_files(scorecard_file=None, base_file=None) -> dict[str, bool]:
    files = get_analysis_files()
    Path(settings.UPLOAD_DIR).mkdir(parents=True, exist_ok=True)
    metadata = load_analysis_file_metadata()
    saved = {
        "scorecard": False,
        "base": False,
    }
    if scorecard_file is not None:
        _write_uploaded_file(files.scorecard_path, scorecard_file)
        metadata["scorecard"] = {
            "original_name": _clean_text(getattr(scorecard_file, "name", "")) or SCORECARD_UPLOAD_NAME,
            "stored_name": files.scorecard_path.name,
            "saved_at": datetime.now().strftime("%d/%m/%Y %H:%M"),
        }
        saved["scorecard"] = True
    if base_file is not None:
        _write_uploaded_file(files.base_path, base_file)
        metadata["base"] = {
            "original_name": _clean_text(getattr(base_file, "name", "")) or BASE_UPLOAD_NAME,
            "stored_name": files.base_path.name,
            "saved_at": datetime.now().strftime("%d/%m/%Y %H:%M"),
        }
        saved["base"] = True
    _save_analysis_file_metadata(metadata)
    return saved


def analysis_files_available() -> bool:
    files = get_analysis_files()
    return files.scorecard_path.exists() and files.base_path.exists()


def load_cached_analysis_snapshot() -> dict[str, object] | None:
    files = get_analysis_files()
    if not files.snapshot_path.exists():
        return None
    with files.snapshot_path.open("r", encoding="utf-8") as file_handle:
        return json.load(file_handle)


def load_analysis_file_metadata() -> dict[str, dict[str, str]]:
    files = get_analysis_files()
    if not files.metadata_path.exists():
        return {}
    with files.metadata_path.open("r", encoding="utf-8") as file_handle:
        return json.load(file_handle)


def build_and_cache_analysis_snapshot() -> dict[str, object]:
    files = get_analysis_files()
    snapshot = build_commercial_analysis(
        files.scorecard_path,
        files.base_path,
        load_analysis_file_metadata(),
    )
    with files.snapshot_path.open("w", encoding="utf-8") as file_handle:
        json.dump(snapshot, file_handle, ensure_ascii=False, indent=2)
    return snapshot


def load_or_build_analysis_snapshot() -> dict[str, object] | None:
    if not analysis_files_available():
        return None
    cached = load_cached_analysis_snapshot()
    if cached:
        return cached
    return build_and_cache_analysis_snapshot()


def build_commercial_analysis(
    scorecard_path: Path,
    base_path: Path,
    file_metadata: dict[str, dict[str, str]] | None = None,
) -> dict[str, object]:
    scorecard_workbook = load_workbook(scorecard_path, data_only=True, read_only=True)
    base_workbook = load_workbook(base_path, data_only=True, read_only=True)
    file_metadata = file_metadata or {}
    try:
        scorecard_summary = _build_scorecard_summary(scorecard_workbook["Scorecard"])
        clients_summary = _build_clients_summary(scorecard_workbook["Clientes"])
        carteira_summary = _build_carteira_summary(scorecard_workbook["Carteira Fundos"])
        base_fundos_summary = _build_base_fundos_summary(base_workbook["Base_Fundos"])
        opening_summary = _build_opening_summary(base_workbook["Fundos Abertos por mês"])
        follow_up_summary = _build_follow_up_summary(scorecard_workbook["BD_FUP"])
        pipeline_summary = _build_pipeline_summary(
            scorecard_workbook[_find_sheet_name(scorecard_workbook.sheetnames, "BD_Neg")],
            scorecard_workbook["BD_Propostas_Vendas"],
        )

        overlap = _build_overlap_summary(
            carteira_summary["funds"],
            base_fundos_summary["funds"],
        )

        return {
            "generated_at": datetime.now().strftime("%d/%m/%Y %H:%M"),
            "source_files": {
                "scorecard": {
                    "original_name": file_metadata.get("scorecard", {}).get("original_name", scorecard_path.name),
                    "stored_name": scorecard_path.name,
                    "saved_at": file_metadata.get("scorecard", {}).get("saved_at", ""),
                },
                "base": {
                    "original_name": file_metadata.get("base", {}).get("original_name", base_path.name),
                    "stored_name": base_path.name,
                    "saved_at": file_metadata.get("base", {}).get("saved_at", ""),
                },
            },
            "reference": {
                "reported_until": scorecard_summary["reported_until"],
                "latest_month_label": scorecard_summary["latest_month_label"],
                "latest_month_date": scorecard_summary["latest_month_date"],
            },
            "metrics": [
                {
                    "label": "Dados atualizados ate",
                    "value": scorecard_summary["reported_until"] or "-",
                    "helper": scorecard_summary["latest_month_label"] or "Sem referencia mensal",
                },
                {
                    "label": "Receita do mes",
                    "value": clients_summary["latest_month_total_formatted"],
                    "helper": clients_summary["latest_month_label"] or "Mes nao identificado",
                },
                {
                    "label": "Atingimento YTD",
                    "value": scorecard_summary["ytd_attainment_label"],
                    "helper": f"Realizado {scorecard_summary['ytd_actual_formatted']} vs target {scorecard_summary['ytd_target_formatted']}",
                },
                {
                    "label": "Fundos cruzados",
                    "value": str(overlap["matched_count"]),
                    "helper": f"{overlap['coverage_on_carteira_label']} da carteira encontrada na base",
                },
                {
                    "label": "Follow-ups mapeados",
                    "value": str(follow_up_summary["total_activities"]),
                    "helper": follow_up_summary["top_activity_label"],
                },
                {
                    "label": "Negocios ganhos",
                    "value": str(pipeline_summary["won_count"]),
                    "helper": f"{pipeline_summary['proposal_count']} propostas enviadas",
                },
            ],
            "kpis": scorecard_summary["kpis"],
            "top_clients": clients_summary["top_clients"],
            "underperforming_clients": clients_summary["underperforming_clients"],
            "carteira_summary": {
                "rows": carteira_summary["row_count"],
                "unique_cnpjs": carteira_summary["unique_cnpjs"],
                "top_groups": carteira_summary["top_groups"],
                "top_managers": carteira_summary["top_managers"],
            },
            "base_summary": {
                "rows": base_fundos_summary["row_count"],
                "unique_cnpjs": base_fundos_summary["unique_cnpjs"],
                "top_managers": base_fundos_summary["top_managers"],
                "top_consultancies": base_fundos_summary["top_consultancies"],
                "opening_by_month": opening_summary,
            },
            "overlap_summary": overlap,
            "follow_up_summary": follow_up_summary,
            "pipeline_summary": pipeline_summary,
        }
    finally:
        scorecard_workbook.close()
        base_workbook.close()


def _build_scorecard_summary(sheet) -> dict[str, object]:
    rows = list(sheet.iter_rows(values_only=True))
    reported_until = _format_date(rows[1][2]) if len(rows) > 1 and len(rows[1]) > 2 else ""
    month_dates = [
        value
        for value in rows[2][3:15]
        if isinstance(value, datetime)
    ] if len(rows) > 2 else []
    latest_month_date = month_dates[-1] if month_dates else None
    latest_month_label = _format_month_year(latest_month_date)

    kpis = []
    current_kpi_name = ""
    current_weight = None
    grouped_rows: list[tuple[str, str, object]] = []
    for row in rows[3:]:
        if not row or not any(value not in (None, "") for value in row[:3]):
            continue
        if row[0]:
            current_kpi_name = str(row[0]).strip()
        if row[1] not in (None, ""):
            current_weight = row[1]
        scenario = str(row[2]).strip() if row[2] not in (None, "") else ""
        if current_kpi_name and scenario:
            grouped_rows.append((current_kpi_name, scenario, row))

    kpi_map: dict[str, dict[str, object]] = {}
    for kpi_name, scenario, row in grouped_rows:
        entry = kpi_map.setdefault(
            kpi_name,
            {
                "name": kpi_name,
                "weight": current_weight,
            },
        )
        actual = row[14] if len(row) > 14 else None
        ytd = row[16] if len(row) > 16 else None
        if scenario.lower().startswith("target"):
            entry["target_value"] = actual
            entry["target_ytd"] = ytd
        elif scenario.lower().startswith("realizado"):
            entry["actual_value"] = actual
            entry["actual_ytd"] = ytd
        elif scenario.lower().startswith("ating."):
            if "ponderado" in scenario.lower():
                entry["weighted_attainment"] = ytd
            else:
                entry["attainment"] = ytd

    for value in kpi_map.values():
        kpis.append(
            {
                "name": value["name"],
                "weight_label": _format_percentage_number(value.get("weight"), digits=1),
                "actual_value": _format_metric_value(value.get("actual_value")),
                "target_value": _format_metric_value(value.get("target_value")),
                "actual_ytd": _format_metric_value(value.get("actual_ytd")),
                "target_ytd": _format_metric_value(value.get("target_ytd")),
                "attainment_label": _format_percentage_number(value.get("attainment")),
                "weighted_attainment_label": _format_percentage_number(value.get("weighted_attainment")),
            }
        )

    revenue_kpi = next((item for item in kpis if item["name"] == "Receita Base"), None)
    realized_ytd = Decimal("0")
    target_ytd = Decimal("0")
    attainment_label = "-"
    if revenue_kpi:
        realized_ytd = _decimal_or_zero(kpi_map["Receita Base"].get("actual_ytd"))
        target_ytd = _decimal_or_zero(kpi_map["Receita Base"].get("target_ytd"))
        attainment_label = _format_percentage_number(
            (realized_ytd / target_ytd) if target_ytd > 0 else Decimal("0")
        )

    return {
        "reported_until": reported_until,
        "latest_month_label": latest_month_label,
        "latest_month_date": latest_month_date.isoformat() if latest_month_date else "",
        "ytd_actual_formatted": format_currency(realized_ytd),
        "ytd_target_formatted": format_currency(target_ytd),
        "ytd_attainment_label": attainment_label,
        "kpis": kpis,
    }


def _build_clients_summary(sheet) -> dict[str, object]:
    rows = list(sheet.iter_rows(values_only=True))
    header_row = rows[2]
    data_rows = [row for row in rows[3:] if row and row[0]]
    month_indexes = [
        index
        for index, value in enumerate(header_row)
        if isinstance(value, datetime)
    ]
    latest_index = month_indexes[13] if len(month_indexes) > 13 else (month_indexes[-1] if month_indexes else None)
    latest_month_label = _format_month_year(header_row[latest_index]) if latest_index is not None else ""

    top_clients = []
    latest_month_total = Decimal("0")
    for row in data_rows:
        month_value = _decimal_or_zero(row[latest_index]) if latest_index is not None else Decimal("0")
        latest_month_total += month_value
        ytd_actual = _decimal_or_zero(row[47]) if len(row) > 47 else Decimal("0")
        ytd_target = _decimal_or_zero(row[48]) if len(row) > 48 else Decimal("0")
        delta = ytd_actual - ytd_target
        top_clients.append(
            {
                "name": str(row[0]).strip(),
                "month_value_raw": month_value,
                "month_value": format_currency(month_value),
                "ytd_actual_raw": ytd_actual,
                "ytd_actual": format_currency(ytd_actual),
                "ytd_target": format_currency(ytd_target),
                "delta_raw": delta,
                "delta": format_currency(delta),
                "coverage": _format_percentage_number((ytd_actual / ytd_target) if ytd_target > 0 else None),
            }
        )

    top_by_month = sorted(top_clients, key=lambda item: item["month_value_raw"], reverse=True)[:10]
    underperforming = [
        item
        for item in sorted(top_clients, key=lambda item: item["delta_raw"])
        if item["delta_raw"] < 0
    ][:10]

    for item in top_by_month + underperforming:
        item.pop("month_value_raw", None)
        item.pop("ytd_actual_raw", None)
        item.pop("delta_raw", None)

    return {
        "latest_month_label": latest_month_label,
        "latest_month_total_formatted": format_currency(latest_month_total),
        "top_clients": top_by_month,
        "underperforming_clients": underperforming,
    }


def _build_carteira_summary(sheet) -> dict[str, object]:
    rows = [
        row for row in list(sheet.iter_rows(values_only=True))[1:]
        if row and any(value not in (None, "") for value in row)
    ]
    funds = []
    group_counter = Counter()
    manager_counter = Counter()
    unique_cnpjs: set[str] = set()
    for row in rows:
        cnpj = _digits_only(row[0])
        unique_cnpjs.add(cnpj)
        manager = _clean_text(row[2])
        group = _clean_text(row[3])
        if group:
            group_counter[group] += 1
        if manager:
            manager_counter[manager] += 1
        funds.append(
            {
                "cnpj": cnpj,
                "fund_name": _clean_text(row[1]),
                "manager_name": manager,
                "client_group": group,
            }
        )
    return {
        "row_count": len(rows),
        "unique_cnpjs": len(unique_cnpjs),
        "funds": funds,
        "top_groups": _counter_rows(group_counter),
        "top_managers": _counter_rows(manager_counter),
    }


def _build_base_fundos_summary(sheet) -> dict[str, object]:
    rows = [
        row for row in list(sheet.iter_rows(values_only=True))[1:]
        if row and any(value not in (None, "") for value in row)
    ]
    funds = []
    manager_counter = Counter()
    consultancy_counter = Counter()
    unique_cnpjs: set[str] = set()
    for row in rows:
        cnpj = _digits_only(row[1])
        unique_cnpjs.add(cnpj)
        manager = _clean_text(row[4])
        consultancy = _clean_text(row[5])
        if manager:
            manager_counter[manager] += 1
        if consultancy:
            consultancy_counter[consultancy] += 1
        funds.append(
            {
                "cnpj": cnpj,
                "fund_name": _clean_text(row[0]),
                "manager_name": manager,
                "consultancy_name": consultancy,
            }
        )
    return {
        "row_count": len(rows),
        "unique_cnpjs": len(unique_cnpjs),
        "funds": funds,
        "top_managers": _counter_rows(manager_counter),
        "top_consultancies": _counter_rows(consultancy_counter),
    }


def _build_opening_summary(sheet) -> list[dict[str, object]]:
    rows = list(sheet.iter_rows(values_only=True))
    summary = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        year_2024 = _decimal_or_zero(row[1]) if len(row) > 1 else Decimal("0")
        year_2025 = _decimal_or_zero(row[2]) if len(row) > 2 else Decimal("0")
        summary.append(
            {
                "month": _clean_text(row[0]),
                "opened_2024": int(year_2024),
                "opened_2025": int(year_2025),
                "delta": int(year_2025 - year_2024),
            }
        )
    return summary


def _build_overlap_summary(carteira_funds: Iterable[dict[str, str]], base_funds: Iterable[dict[str, str]]) -> dict[str, object]:
    carteira_by_cnpj = {
        fund["cnpj"]: fund
        for fund in carteira_funds
        if fund.get("cnpj")
    }
    base_by_cnpj = {
        fund["cnpj"]: fund
        for fund in base_funds
        if fund.get("cnpj")
    }
    matched_cnpjs = set(carteira_by_cnpj) & set(base_by_cnpj)
    only_carteira = [
        carteira_by_cnpj[cnpj]
        for cnpj in sorted(set(carteira_by_cnpj) - matched_cnpjs)
    ][:10]
    only_base = [
        base_by_cnpj[cnpj]
        for cnpj in sorted(set(base_by_cnpj) - matched_cnpjs)
    ][:10]
    carteira_count = len(carteira_by_cnpj)
    coverage = (Decimal(len(matched_cnpjs)) / Decimal(carteira_count)) if carteira_count else Decimal("0")
    return {
        "matched_count": len(matched_cnpjs),
        "only_carteira_count": len(set(carteira_by_cnpj) - matched_cnpjs),
        "only_base_count": len(set(base_by_cnpj) - matched_cnpjs),
        "coverage_on_carteira_label": _format_percentage_number(coverage),
        "only_carteira": [
            {
                "fund_name": item.get("fund_name", ""),
                "manager_name": item.get("manager_name", ""),
                "client_group": item.get("client_group", ""),
                "cnpj": _format_cnpj(item.get("cnpj", "")),
            }
            for item in only_carteira
        ],
        "only_base": [
            {
                "fund_name": item.get("fund_name", ""),
                "manager_name": item.get("manager_name", ""),
                "consultancy_name": item.get("consultancy_name", ""),
                "cnpj": _format_cnpj(item.get("cnpj", "")),
            }
            for item in only_base
        ],
    }


def _build_follow_up_summary(sheet) -> dict[str, object]:
    rows = list(sheet.iter_rows(values_only=True))[1:]
    activity_counter = Counter()
    company_counter = Counter()
    for row in rows:
        if not row:
            continue
        activity_type = _clean_text(row[1])
        company_name = _clean_text(row[8])
        if activity_type:
            activity_counter[activity_type] += 1
        if company_name:
            company_counter[company_name] += 1
    top_activity = activity_counter.most_common(1)
    top_activity_label = (
        f"Tipo dominante: {top_activity[0][0]} ({top_activity[0][1]})"
        if top_activity
        else "Sem atividades mapeadas"
    )
    return {
        "total_activities": sum(activity_counter.values()),
        "top_activity_label": top_activity_label,
        "activity_types": _counter_rows(activity_counter),
        "top_companies": _counter_rows(company_counter),
    }


def _build_pipeline_summary(negocios_sheet, propostas_sheet) -> dict[str, object]:
    negocio_rows = list(negocios_sheet.iter_rows(values_only=True))[1:]
    proposal_rows = list(propostas_sheet.iter_rows(values_only=True))[1:]
    stage_counter = Counter()
    company_counter = Counter()
    source_counter = Counter()
    for row in negocio_rows:
        if not row:
            continue
        stage = _clean_text(row[3])
        company = _clean_text(row[4])
        if stage:
            stage_counter[stage] += 1
        if company:
            company_counter[company] += 1
    for row in proposal_rows:
        if not row:
            continue
        source = _clean_text(row[11])
        if source:
            source_counter[source] += 1
    return {
        "won_count": stage_counter.get("Venda Ganha", 0),
        "proposal_count": stage_counter.get("Proposta Enviada", 0),
        "stages": _counter_rows(stage_counter),
        "top_companies": _counter_rows(company_counter),
        "proposal_sources": _counter_rows(source_counter),
    }


def _counter_rows(counter: Counter[str], limit: int = 10) -> list[dict[str, object]]:
    return [
        {
            "label": label,
            "count": count,
        }
        for label, count in counter.most_common(limit)
    ]


def _write_uploaded_file(destination: Path, uploaded_file) -> None:
    with destination.open("wb+") as file_handle:
        for chunk in uploaded_file.chunks():
            file_handle.write(chunk)


def _save_analysis_file_metadata(metadata: dict[str, dict[str, str]]) -> None:
    files = get_analysis_files()
    with files.metadata_path.open("w", encoding="utf-8") as file_handle:
        json.dump(metadata, file_handle, ensure_ascii=False, indent=2)


def _find_sheet_name(sheet_names: Iterable[str], prefix: str) -> str:
    for name in sheet_names:
        if str(name).startswith(prefix):
            return str(name)
    raise KeyError(f"Aba com prefixo {prefix} nao encontrada.")


def _decimal_or_zero(value) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value))
    except Exception:
        return Decimal("0")


def _format_metric_value(value) -> str:
    decimal_value = _decimal_or_zero(value)
    if decimal_value == decimal_value.to_integral_value():
        return format_currency(decimal_value)
    return _format_percentage_number(decimal_value)


def _format_percentage_number(value, digits: int = 1) -> str:
    if value in (None, ""):
        return "-"
    decimal_value = _decimal_or_zero(value) * Decimal("100")
    quantizer = Decimal("1") if digits == 0 else Decimal(f"1.{'0' * digits}")
    formatted = f"{decimal_value.quantize(quantizer):,.{digits}f}"
    formatted = formatted.replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{formatted}%"


def _format_month_year(value) -> str:
    if not isinstance(value, datetime):
        return ""
    month_names = [
        "jan",
        "fev",
        "mar",
        "abr",
        "mai",
        "jun",
        "jul",
        "ago",
        "set",
        "out",
        "nov",
        "dez",
    ]
    return f"{month_names[value.month - 1]}/{str(value.year)[-2:]}"


def _format_date(value) -> str:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    return _clean_text(value)


def _clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _digits_only(value) -> str:
    return "".join(char for char in _clean_text(value) if char.isdigit())


def _format_cnpj(value: str) -> str:
    digits = _digits_only(value)
    if len(digits) != 14:
        return digits
    return f"{digits[:2]}.{digits[2:5]}.{digits[5:8]}/{digits[8:12]}-{digits[12:]}"
